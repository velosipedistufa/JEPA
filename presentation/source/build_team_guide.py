"""Team guide for the JEPA defence: concepts in plain words, then every slide with what to say.
python3 build_team_guide.py <dir with slide-NN.jpg renders of JEPA_defence_visual.pptx>
(needs python-docx on PYTHONPATH)
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent
THUMBS = Path(sys.argv[1])
NAVY, TEAL, MUTED = RGBColor(0x14, 0x2A, 0x3B), RGBColor(0x06, 0x7F, 0x7C), RGBColor(0x61, 0x71, 0x7B)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)  # A4
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Inches(0.9))
sec.top_margin = sec.bottom_margin = Inches(0.8)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.12
for name, size, color in (("Title", 26, NAVY), ("Heading 1", 18, NAVY), ("Heading 2", 14, TEAL), ("Heading 3", 12, NAVY)):
    s = doc.styles[name]
    s.font.name = "Calibri"
    s.font.size = Pt(size)
    s.font.color.rgb = color
    s.font.bold = True
    rpr = s.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(a), "Calibri")
    for el in rpr.findall(qn("w:i")):
        rpr.remove(el)


def para(text="", bold_lead=None, italic=False, color=None, size=None, after=None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    r = p.add_run(text)
    r.italic = italic
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            p.add_run(it[0]).bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)
        p.paragraph_format.space_after = Pt(3)


def label(text):
    p = doc.add_paragraph()
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = TEAL
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True


def say(text):
    """Narration to practise, set off in a shaded box."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), "E8F3F1")
    c._tc.get_or_add_tcPr().append(sh)
    c.paragraphs[0].add_run(text).font.size = Pt(11)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def table(rows, widths):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = t.cell(i, j)
            cell.width = Inches(widths[j])
            cell.text = ""
            r = cell.paragraphs[0].add_run(v)
            r.font.size = Pt(10)
            r.bold = i == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------------ cover
doc.add_paragraph("Language over demo breadcrumbs", style="Title")
para("Team guide to the defence deck: the ideas in plain words, then every slide with what to say", color=MUTED, size=13)
para("Alexander Bashirov · Didar Seyidov · Faith Ogunwoye · Skoltech", color=MUTED)
para(
    "How to use this guide. Part 1 tells the whole story in one minute. Part 2 explains every term the deck uses, "
    "assuming no background. Part 3 walks through the 17 slides: what each one shows, a narration you can practise "
    "out loud (the shaded boxes), the numbers worth remembering, and the questions it may trigger. Part 4 collects "
    "likely panel questions with short answers. Every number here comes from the saved runs in the JEPA repo "
    "(branch 5090-run-demo-breadcrumbs, folder research/artifacts/5090)."
)

doc.add_heading("Part 1 · The project in one minute", level=1)
para(
    "A robot needs to imagine what its actions will do before it acts. LeWorldModel (LeWM) is a learned "
    "“imagination” for a simple pushing game called Push-T: it turns a camera image into 192 numbers and "
    "predicts how those numbers change when the robot moves. A planner tries hundreds of action sequences inside "
    "that imagination and picks the one that ends closest to a goal picture."
)
para(
    "The catch: this works well when the goal is close (about 25 steps away, 47 of 50 tries succeed) and badly "
    "when it is far. Our course question was whether a sentence, such as “turn the T clockwise”, can name a "
    "goal this planner can actually reach."
)
para(
    "Our answer has two parts. First, we never ask the planner for a far goal. We break the journey into short hops "
    "using real expert demonstrations as breadcrumbs: a map of 427,396 demo states, linked across different demos, "
    "gives a route; the planner chases the next breadcrumb about 25 steps ahead, then looks again. Second, a small "
    "local language model (Qwen2.5-1.5B) turns a free sentence into a goal region the map understands. Breadcrumbs "
    "lift success from 10/50 to 29/50 on goals 100 steps away; free sentences work end to end in 51 of 111 tasks. "
    "We also tried feeding the words straight into the model’s latent space; that works for “put it on the "
    "target” but not yet for turns and pushes."
)

# ------------------------------------------------------------------ glossary
page_break()
doc.add_heading("Part 2 · Key terms, explained from scratch", level=1)
para("Read this once before practising. The deck uses these words without stopping to define them.", italic=True, color=MUTED)

doc.add_heading("The model", level=2)
for term, text in [
    ("World model. ", "A network that predicts what happens next if the robot takes some action. With it, a robot can try "
     "actions in its head before trying them for real. Think of it as the robot’s imagination."),
    ("Latent (embedding, z). ", "A short list of numbers that summarises an image. LeWM squeezes a 224×224 camera "
     "image into 192 numbers. Images that look alike end up with similar numbers. “Latent space” is the space "
     "of all such number lists. Comparing two latents is much cheaper than comparing two images pixel by pixel."),
    ("JEPA (Joint-Embedding Predictive Architecture). ", "An idea proposed by Yann LeCun in 2022 [1]. Instead of "
     "predicting the next image pixel by pixel (hard and wasteful, because most pixels are irrelevant detail), a JEPA "
     "predicts the next latent. Both the current and the future image are encoded, and the network learns to predict "
     "one embedding from the other. I-JEPA [2] applied this to still images; LeWM applies it to video with actions."),
    ("LeWorldModel (LeWM). ", "A JEPA world model trained end to end from pixels (Maes et al. 2026 [3]). It has two "
     "parts: an encoder (a small Vision Transformer, ViT-tiny) that maps an image to a 192-number latent, and a "
     "predictor that takes the current latent and a block of actions and predicts the next latent. We use the public "
     "Push-T checkpoint (quentinll/lewm-pusht) and never retrain it. That is what “frozen” means in the deck."),
    ("Frozen. ", "The weights are not changed. Everything we add sits around the model. This matters because "
     "training a world model needs far more compute than a course project has, and it makes the comparison fair: "
     "any gain comes from what we feed the planner, not from a better model."),
]:
    para(text, bold_lead=term)

doc.add_heading("The task", level=2)
for term, text in [
    ("Push-T. ", "A standard robot-learning benchmark [5]. Seen from above, a grey T-shaped block lies on a 512×512 "
     "pixel table. A small round pusher (the “agent”) moves around and pushes it. A green T outline marks the "
     "target pose. The full state is just five numbers: pusher x and y, block x and y, block angle."),
    ("Success. ", "The simulator’s own rule: the block and pusher are within 20 pixels of the goal and the block angle "
     "is within 20°. For word commands we score the block only, with the same 20 px / 20° tolerances."),
    ("Expert demonstrations (demos). ", "Recordings of a skilled operator solving Push-T: 18,685 episodes, about 2.3 "
     "million frames. We use 90% of the episodes (16,817) to build the map and keep 10% unseen for testing."),
    ("Step, budget. ", "One step is one small pusher move. The budget is how many steps an attempt may use; all our "
     "scored runs use 300."),
    ("Offset (+50, +100). ", "For demo tasks we start at a frame of an unseen demo and use the frame 50 or 100 steps "
     "later in the same demo as the goal. Bigger offset = farther goal."),
    ("Live goal. ", "A random goal state drawn by the simulator, which no demo needs to have visited. Much harder."),
]:
    para(text, bold_lead=term)

doc.add_heading("The planner", level=2)
for term, text in [
    ("Planning with a world model. ", "Try many candidate action sequences in imagination, score how close each one "
     "ends to the goal, and execute the best one. No policy network is trained."),
    ("Cost. ", "The score of a candidate: the squared distance between the latent the model imagines at the end and "
     "the latent of the goal image. Lower is better."),
    ("CEM (cross-entropy method). ", "The search rule [6]. Sample 300 random action sequences, imagine each, keep the "
     "best 30, fit a new sampling distribution to those 30, and repeat 30 times. The final average is the plan. Each "
     "plan covers 25 steps (5 blocks of 5 actions)."),
    ("Replanning (receding horizon). ", "Execute the plan, look at the new camera image, plan again. Errors get "
     "corrected every 25 steps."),
    ("Why far goals fail. ", "The planner only looks 25 steps ahead. If the goal is 100 steps away, no 25-step plan "
     "gets close, and latent distance is a poor compass that far out: moving toward the goal can look worse in latent "
     "space before it looks better. The planner then wanders or stalls. Turning the T in place is similar: the image "
     "changes little until the turn is nearly done."),
    ("Flat planner. ", "The baseline: the official planner chasing the final goal image directly, the whole time."),
]:
    para(text, bold_lead=term)

doc.add_heading("Our method", level=2)
for term, text in [
    ("Subgoal / breadcrumb. ", "An intermediate target. Instead of the far goal, the planner chases a nearby frame "
     "that lies on the way. The image is of a learner driver who can only drive 25 metres at a time: you show them "
     "the next stretch from an expert’s dashcam, not the destination."),
    ("Demo graph (the map). ", "One node every 5 steps of every training demo: 427,396 nodes. Two kinds of links: "
     "“next frame” links follow each demo forward (410,579), and jump links join states from different demos "
     "that are close in state (2,852,207). Closeness mixes pusher position, block position and block angle; jump "
     "links need a distance under 15 in those units."),
    ("Route. ", "The cheapest path through the graph from the node nearest the current state to a goal node, found "
     "with Dijkstra’s shortest-path algorithm. A route can start in one demo and finish in another."),
    ("Hop. ", "One cycle: find where we are on the map, walk about 25 demo steps along the route, render that demo "
     "state as an image, encode it, plan toward it with CEM for 25 steps, execute, repeat. When the route reaches "
     "the goal, chase the goal itself."),
    ("Retrieval by state. ", "To find “where we are” on the map we use the simulator’s true state. On a "
     "real robot, a marker on the block would give the same numbers. LeWM itself only ever sees images."),
    ("Hi-LeWM. ", "The closest prior work (Caselli et al. 2026 [4]). It showed that later frames of the same demo "
     "make good subgoals. That needs the demo the start came from. Our graph finds routes through other demos, so "
     "it also works from starts no demo began at."),
    ("Goal region. ", "A command names a set of acceptable poses, not one picture. “Put it on the target”: "
     "block within 20 px and 20° of the green outline. “Turn 90° clockwise”: angle changed by "
     "+90° (±20°), block within 40 px of where it started. “Push left 80 px”: block within "
     "20 px of the point 80 px to the left, any angle. Every graph node inside the region is a goal node."),
    ("Clockwise, up. ", "As seen on screen. The screen’s y axis points down, so “up” means smaller y, and "
     "a positive block angle is a clockwise turn. We checked this on the renderer before running."),
]:
    para(text, bold_lead=term)

doc.add_heading("Language", level=2)
for term, text in [
    ("LLM parser. ", "Qwen2.5-1.5B-Instruct [7], a small open language model running locally on the GPU (about 0.14 s "
     "per sentence). It reads a sentence and writes a structured command, for example {\"action\": \"turn\", "
     "\"rotation\": \"clockwise\", \"deg\": 90}. It learns the format from ten examples in its prompt "
     "(“few-shot”); it is not fine-tuned. Greedy decoding means it always picks its most likely answer."),
    ("Keyword baseline. ", "Simple rules built from the seven template sentences (look for “target”, "
     "“clockwise”, “left”…). It shows how much the language model adds: 31/120 versus 101/120."),
    ("Held-out paraphrases. ", "120 test sentences written by a different model (Codex) that never saw the parser. "
     "They include casual, indirect and deliberately odd phrasings (typos, slang, “that tee”). They were "
     "committed to git before the parser was scored, so the parser could not be tuned to them."),
    ("Gold command. ", "The correct command for a test sentence (written with the sentence). “Gold” runs show "
     "what perfect parsing would achieve; “parsed” runs use what Qwen actually produced. The gap is the cost "
     "of the language layer."),
    ("Subgoal head (“words into the latent”). ", "A small network that takes the start latent, the current "
     "latent and a sentence embedding, and outputs the latent of the next breadcrumb. It was trained on 22,420 "
     "examples cut from graph routes. At test time there is no graph and no parser: the words steer the planner "
     "through latent space directly."),
    ("Sentence embedding. ", "The sentence turned into numbers by the same frozen Qwen model (the average of its last "
     "layer). Similar meanings give similar vectors."),
    ("MSE and the copy baseline. ", "Mean squared error between the predicted and the true next latent. The copy "
     "baseline predicts “nothing changes” (the current latent). 0.67× means the head’s error is "
     "a third smaller than doing nothing."),
    ("Shuffled control. ", "The same head, but every sentence is swapped with another row’s sentence. If results "
     "barely drop, the words were not doing the work. They drop from 18 to 10 of 111."),
]:
    para(text, bold_lead=term)

doc.add_heading("How we report results", level=2)
for term, text in [
    ("Pre-registration: claim and kill line. ", "Before each scored run we committed to git what we expected (the "
     "claim) and the result that would kill the idea (the kill line). Constants were not changed after seeing a "
     "score. “Holds” = claim met. “Between” = claim missed but kill line not reached."),
    ("Paired outcomes. ", "Flat and breadcrumbs run on exactly the same tasks, so we can count tasks only one method "
     "solved. At offset 100: 23 solved only with breadcrumbs, 4 only flat."),
    ("pp (percentage points). ", "The difference between two percentages: 58% versus 20% is +38 pp."),
    ("Strict in-place gate. ", "A stricter turn test from an earlier experiment (artifact 012): block drift under "
     "15 px. Our turn task allowed 40 px; the deck reports both."),
]:
    para(text, bold_lead=term)

# ------------------------------------------------------------------ slides
SLIDES = [
    ("Title", "Title, team, and the one-line idea: words name the goal, demonstrations supply the steps. The picture is a real logged Push-T start state.",
     "Hello. We are Alexander, Didar and Faith. Our project asks whether a frozen world model, one we never retrain, can follow spoken-style instructions. Our one-line answer: let words name the goal, and let expert demonstrations supply the steps to get there. On the right is the Push-T task we work on: push the grey T onto the green outline.",
     ["Model frozen throughout", "Task: Push-T (pusher + T-shaped block)"],
     []),
    ("What is new here", "The four contributions, as a promise of the talk: words name the goal; routes across many demos; short hops reach far goals; words straight into the latent.",
     "Here is what this project adds; each point comes back later with numbers. One: a sentence becomes a goal region, not a single goal image, and a small local language model does the reading. Two: we build a map of expert demonstrations that links different demos, so we can find a route even from a start no demo began at. Prior work used later frames of the same demo. Three: the planner is only ever asked to reach something about 25 steps away, which is where it is reliable. Four: a first test of feeding the words straight into the model’s latent space. And every claim was written down in git before we ran it.",
     ["101/120 held-out sentences parsed", "427,396 demo states in the map", "10/50 → 29/50 at 100 steps", "18/111 vs 10/111 with shuffled words"],
     ["“Isn’t this just Hi-LeWM?” No: Hi-LeWM uses the same demo the start came from. We route through other demos, which a new situation needs."]),
    ("01 · The problem", "Three cards: short demo goal 47/50; long live goal 0/20; in-place turn 0/20. These are different settings, not a controlled sweep.",
     "The frozen planner is strong over a short hop: 47 of 50 goals 25 steps ahead, matching the 45 of 50 our teammate got on a different GPU. But give it a random far goal and it gets 0 of 20; ask it to turn the T in place and it gets 0 of 20. So the course question became: can a sentence name a goal this existing planner can actually follow?",
     ["47/50 short goals (5090), 45/50 on the 4090", "0/20 random live goals at budget 50", "0/20 rendered in-place turns"],
     ["“Are these comparable?” No, they are three different settings, shown to motivate the problem, not as one experiment."]),
    ("02 · What we tried first", "Three diagnostics: the angle can be read from the latent (12–15° error); a fixed latent step per command does not fit (residual 0.991); planning a turn directly fails (2/20, strict 0/20; three rendered stages 1/20).",
     "Before the breadcrumbs, we checked whether the problem was the model’s understanding. It is not: the block angle can be read from the latent to within about 12 to 15 degrees. But a fixed “turn” direction in latent space does not exist, and planning a turn directly almost never works, even with three intermediate rendered goals. So the model knows the angle; the planner just cannot get there. Decision: keep the model frozen and change what we give the planner as a subgoal.",
     ["Angle error 12.08° (180° frames), 14.58° (near 0°)", "Residual ratio 0.991 (no single Δz per command)", "Angle-cost turn 2/20, strict 0/20, staged 1/20"],
     ["“What is residual ratio?” Error of the best constant-step fit divided by the size of the target: 0.991 means the fit explains almost nothing."]),
    ("03 · The idea", "The pipeline left to right: sentence → goal region; graph route through demos (schematic); plan, act, repeat on the next demo frame.",
     "Here is the method. A sentence like “turn the T clockwise” is parsed into a goal region. The demo map then finds a route of real demo states from where the block is to that region, switching between demos where they pass close by. We render the demo state about 25 steps along the route, encode it with the frozen model, let the unchanged planner chase it for 25 steps, then look again. Long task, short hops.",
     ["427,396 nodes, one every 5 steps, training demos only", "Hop length about 25 steps", "Retrieval uses the simulator state"],
     ["“Isn’t using the simulator state cheating?” It is only used to find our place on the map; the planner and model only see images. On a robot a tag on the block gives that state.", "The graph picture is a schematic, not a logged route."]),
    ("04 · Demo breadcrumbs", "Bar chart, flat versus breadcrumbs on four task sets, same planner and budget. Callout: +38 pp at offset 100; 23 gained, 4 lost.",
     "This is the main result. Same planner, same budget; only the subgoals change. Goals 50 steps away: 48% to 68%. Goals 100 steps away: 20% to 58%, a gain of 38 points; the farther the goal, the bigger the help. On word-template tasks: 28% to 53%. Random live goals stay hard, 0 to 15%. Looking task by task at 100 steps, breadcrumbs won 23 tasks flat lost, and lost only 4.",
     ["+50: 24/50 → 34/50", "+100: 10/50 → 29/50 (+38 pp)", "Live: 0/20 → 3/20", "Templates: 18/64 → 34/64"],
     ["“Is it significant?” At 100 steps the paired split is 23 vs 4, a clear effect; at 50 steps 16 vs 6 is promising but borderline. We report counts, not p-values.", "“Why are live goals poor?” They are far from every demo state (Backup D)."]),
    ("05 · Words as the goal", "Left: 101/120 sentences parsed correctly (keyword rules 31/120). Right: end to end, gold command 57/111, parsed command 51/111.",
     "Now free language. We tested 120 sentences written by a different model, including odd ones with typos and slang. The small language model parses 101 of them correctly, where keyword rules manage 31. Running the robot end to end, perfect commands succeed 57 times out of 111 and the parsed ones 51 times. So reading the sentence costs only 6 successes; most failures come from execution, not language.",
     ["Parsing 84.2% (claim was 90%, kill 75%)", "Keyword rules 25.8%", "Gold 57/111, parsed 51/111, −5.4 pp", "9 pushes excluded (target outside the table)"],
     ["“Why a 1.5B model?” Small, local, fast (0.14 s per sentence), no API; good enough to show the interface.", "“Who wrote the test sentences?” Codex, a different model, before scoring."]),
    ("06 · Turns are not in-place", "Scatter of block drift for all 20 turn tasks. 8/20 turned within 40 px; none stayed within 15 px.",
     "We want to be precise about turns. With breadcrumbs, 8 of 20 turn commands reach the right angle while the block stays within 40 pixels. But the T slides while it turns: none of them pass the strict 15-pixel in-place test. So we do not claim in-place rotation. Interestingly, the demos do contain in-place turns, in about 10% of episodes, so better routing could help.",
     ["Breadcrumbs 8/20 turns, strict 0/20", "Flat 4/20, strict 1/20", "In-place turns in demos: 8,040 nodes, 1,691 episodes"],
     ["“Why does it drift?” Expert demos usually turn the T by pushing it along; routes inherit that."]),
    ("07 · A head instead of the graph", "The learned subgoal head: 0.67× validation error versus copying, but 18/111 end to end, versus 51–57 with the graph; 10/111 with shuffled words.",
     "Our supervisor’s question was: can the words go into the model itself? We trained a small head that turns the current latent and the sentence into the next subgoal latent, with no graph and no parser at test time. Offline it is clearly better than guessing nothing changes. On the robot it reaches 18 of 111, against 10 with the wrong sentence, so the words do steer it, but far below the graph’s 51. Lower prediction error did not become better control.",
     ["Val MSE 0.67× copy baseline", "Head 18/111, shuffled 10/111", "Graph + parser 51/111, graph + gold 57/111"],
     ["“Why so much worse?” Turns and pushes are relative to the start pose; with 22k examples the head cannot imagine those poses well. It works mainly for “put it on the target” (Backup B)."]),
    ("08 · What the defence should take", "Three takeaways: nearby demo frames make distant goals easier; language can select the route; reachability is still the limit.",
     "Three things to take away. Nearby demo frames make distant goals much easier for the same frozen planner. Language can select the route, with only a small loss from parsing. And reachability is still the limit: strict in-place turns and the learned head show where the planner cannot yet go.",
     ["10/50 → 29/50", "101/120; 57 → 51", "Strict turns 0/20; head 18/111"],
     []),
    ("09 · Conclusion", "What we showed and what comes next, then thank you.",
     "To conclude: a frozen world model can follow words, one reachable hop at a time. Breadcrumbs from expert demos lift success by 20 to 38 points; free sentences work end to end; and words can enter the latent directly, with a first signal. Next we want strict in-place turns, better coverage for random goals, a stronger latent head, and finally the real lab UR3 arm with a camera, a 3D-printed T and teleoperated demos. Thank you; we are happy to take questions.",
     ["+20 to +38 pp on demo goals", "51/111 end to end (57 with perfect parsing)"],
     ["“What would the robot version need?” Teleoperated demos on the table, a tag on the T for the state, and LeWM retrained on that camera."]),
    ("Backup A · Paired evidence", "For each task set: tasks both methods solved, flat only, breadcrumbs only, neither.",
     "Use when asked whether the gain is real. Because both methods ran on identical tasks we can count wins and losses directly: at 100 steps breadcrumbs gained 23 tasks and lost 4.",
     ["+50: 16 vs 6", "+100: 23 vs 4", "Live: 3 vs 0", "Templates: 24 vs 8"], []),
    ("Backup B · Action breakdown", "Graph + parser, head, and shuffled head, split by command type.",
     "Use when asked what the latent head learned. Its advantage over shuffled words is almost entirely on “put it on the target”: 11 of 30 against 2. On turns and pushes it is no better than the wrong sentence.",
     ["Target: graph 20/30, head 11/30, shuffled 2/30", "Turn: 17/45, 3/45, 3/45", "Push: 14/36, 4/36, 5/36"], []),
    ("Backup C · A measured example", "One real clockwise-turn task: start and recorded end, reconstructed from the logged block poses.",
     "Use to make the drift concrete. This task counted as a success: the T turned 77 degrees clockwise, within 20 of the 90 asked, but it moved 22.5 pixels, so it fails the strict in-place gate.",
     ["Turn 77.2° (asked +90°)", "Drift 22.5 px"], ["The images are reconstructed endpoints, not a video; the pusher is omitted."]),
    ("Backup D · The demo graph", "Graph size and coverage: demo goals 50/50 with a route, word commands 56/64, random live goals 13/20 but far from any demo state.",
     "Use when asked why live goals are hard. Every held-out demo goal has a route, and the start is close to a demo state. Random live goals sit about four times farther from any demo state, so the first and last hops are long and the planner leaves its comfort zone.",
     ["Nodes 427,396; next-frame links 410,579; jump links 2,852,207", "Gap to nearest demo state: 9 (demo) vs 37–39 (live)"], []),
    ("Backup E · Every check", "Table of all twelve pre-registered checks with claim, kill line, result and verdict.",
     "Use when asked how honest the numbers are. Each line was committed before its run. Seven hold, five land between claim and kill, none were killed.",
     ["7 hold, 5 between, 0 killed"], []),
    ("References", "The nine sources and the repository link.",
     "No narration needed; leave it up during questions.", [], []),
]

page_break()
doc.add_heading("Part 3 · The slides, one by one", level=1)
para("About 1.5 minutes per main slide; slides 1–11 make a 14–16 minute talk. Backups A–E and the references are for questions.", italic=True, color=MUTED)
for i, (title, shows, narration, numbers, asks) in enumerate(SLIDES, start=1):
    doc.add_heading(f"Slide {i} · {title}", level=2)
    img = THUMBS / f"slide-{i:02d}.jpg"
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    if img.exists():
        doc.add_picture(str(img), width=Inches(4.6))
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    label("What it shows")
    para(shows)
    label("What to say")
    say(narration)
    if numbers:
        label("Numbers to know")
        bullets(numbers)
    if asks:
        label("If asked")
        bullets(asks)

# ------------------------------------------------------------------ Q&A
page_break()
doc.add_heading("Part 4 · Likely questions", level=1)
for q, a in [
    ("Why not train the world model with language?", "Retraining LeWM needs far more compute than we have, and it would change two things at once. Keeping it frozen shows the gain comes from the subgoal interface alone."),
    ("Is using simulator state for the map cheating?", "The model and planner only see images. The state is used to find our place on the map, exactly what a marker on the block would give a real robot. A latent-only version is future work."),
    ("What is new compared with Hi-LeWM?", "Hi-LeWM uses later frames of the same demo, which requires the demo the start came from. We route through many different demos, so it works from new starts, and we add language on top."),
    ("Why does the planner fail on far goals at all?", "It looks only 25 steps ahead, and latent distance is a poor compass far away. Breadcrumbs keep every request inside that 25-step range."),
    ("Are the gains statistically solid?", "At 100 steps, 23 tasks gained versus 4 lost on identical tasks is a clear effect. At 50 steps (16 vs 6) it is promising but borderline. We report descriptive counts and do not claim significance."),
    ("Why did some claims miss?", "We set the bars before running and kept them. Parsing reached 84% against a 90% claim; parsed execution lost 5.4 points against a 5-point claim. None hit a kill line."),
    ("Why do turns drift?", "The demos usually turn the T while pushing it along, and the route follows the demos. Strict in-place turns exist in about 10% of demos; routing toward them is a next step."),
    ("Could a bigger language model help?", "Parsing is not the main loss: perfect commands give 57/111 and parsed ones 51/111. Execution is the bigger gap."),
    ("What did the latent head show?", "Words in the latent can steer the planner (18 vs 10 with shuffled words), mainly for “put it on the target”. For relative commands it needs more data or a better design. The graph route stays about three times better."),
    ("What is the real-robot plan?", "The lab UR3 arm with a camera over a table, a 3D-printed T with a tag, teleoperated demos, and LeWM retrained on that camera. The code avoids simulator-only shortcuts where possible so it carries over."),
]:
    para(a, bold_lead=q + " ")

doc.add_heading("Part 5 · Timing and hand-offs", level=1)
bullets([
    ("Slides 1–3 (about 3 min): ", "who we are, what is new, the problem."),
    ("Slides 4–6 (about 4.5 min): ", "what we tried first, the idea, the main breadcrumb result."),
    ("Slides 7–9 (about 4.5 min): ", "language, turns, the latent head."),
    ("Slides 10–11 (about 2 min): ", "takeaways and conclusion."),
    ("Questions: ", "jump to Backup A (gains), B (head), C (turn example), D (graph coverage), E (all checks)."),
])

doc.add_heading("Sources", level=1)
for r in [
    "[1] LeCun, Y. (2022). A Path Towards Autonomous Machine Intelligence. OpenReview.",
    "[2] Assran, M., et al. (2023). Self-Supervised Learning from Images with a Joint-Embedding Predictive Architecture. CVPR.",
    "[3] Maes, L., Le Lidec, Q., Scieur, D., LeCun, Y., Balestriero, R. (2026). LeWorldModel. arXiv:2603.19312.",
    "[4] Caselli, Massafra, Punzo, Lo Sardo, Pantelidis, Bhethanabhotla (2026). Hi-LeWM. arXiv:2607.12547.",
    "[5] Florence, P., et al. (2021). Implicit Behavioral Cloning. CoRL.",
    "[6] Rubinstein, R. Y. (1999). The Cross-Entropy Method for Combinatorial and Continuous Optimization.",
    "[7] Yang, A., et al. (2024). Qwen2 Technical Report. arXiv:2407.10671.",
    "Repository: github.com/velosipedistufa/JEPA, branch 5090-run-demo-breadcrumbs (HANDOFF_5090.md, research/artifacts/5090).",
]:
    para(r, size=10)

path = OUT / "JEPA_defence_team_guide.docx"
doc.save(path)
print(path)
